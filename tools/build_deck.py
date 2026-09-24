#!/usr/bin/env python3
"""Construit src/data/deck.json depuis le classeur Excel du vocabulaire.

Usage :
    python3 tools/build_deck.py docs/deck-source.xlsx src/data/deck.json

Le classeur a quatre feuilles, dans deux formats :
  - Sheet1 : une colonne "terme;définition;exemple;#tag" (export Anki), avec
    parfois des trous {{c1::...}} dans la définition ;
  - V2, Code, Code V2 : deux colonnes "terme" | "définition<br><br><b>Example :</b><br>exemple".

Sheet1 fait foi quand un terme existe dans plusieurs feuilles : c'est la seule
qui porte un tag. Le script est déterministe : même classeur → même JSON.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, OrderedDict

import openpyxl

CLOZE = re.compile(r"\{\{c\d+::(.*?)(?:::.*?)?\}\}")
TAGS = re.compile(r"<[^>]+>")

# Matières de l'app. Le sujet "web" = HTML / CSS / JavaScript.
SUBJECT_BY_TAG_PREFIX = {
    "#AI_": "ia",
    "#Code_": "web",
    "#Data": "ia",
}

# Sous-thèmes Python attribués par listes de termes (la feuille Code n'a pas de tag).
PYTHON_TOPICS = {
    "terminal": {"pwd", "ls", "cd", "mkdir", "touch", "cp", "mv", "rm", "clear", "sudo",
                 "pip install", "python", "cat", "grep", "history"},
    "pandas": None,   # tout ce qui commence par pd. / df. / import pandas
    "requests": {"requests.get()", "requests.post()", "response.status_code",
                 "response.json()", "headers={}", "params={}", "time.sleep()"},
    "fichiers": {"open()", "with open(...) as f:", ".write()", "os.listdir()"},
    "operateurs": {"vars()", "x + y", "x - y", "x * y", "x / y", "x // y", "x % y", "bool"},
    "erreurs": None,  # tout ce qui finit par Error
}


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return text or "x"


def clean(text: str) -> str:
    text = text.replace("<br>", "\n").replace("<br/>", "\n")
    text = TAGS.sub("", text)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&").replace("&quot;", '"')
    return re.sub(r"[ \t]+", " ", text).strip()


def split_v2(cell: str) -> tuple[str, str]:
    """'définition<br><br><b>Example :</b><br>exemple' → (définition, exemple)."""
    parts = re.split(r"<b>\s*Example\s*:?\s*</b>", cell, maxsplit=1, flags=re.I)
    definition = clean(parts[0])
    example = clean(parts[1]) if len(parts) > 1 else ""
    return definition, example


def split_sheet1(middle: list[str]) -> tuple[str, str]:
    """Colonnes entre le terme et le tag → (définition, exemple).

    Un ';' en trop vient d'une définition ou d'un exemple qui en contient un.
    Quand une colonne finit une phrase (« …développe.;GPT-4 est closed
    source;tu ne peux… »), la définition s'arrête là et la suite rejoint
    l'exemple : sinon « .;GPT-4… » s'affichait dans la définition.
    """
    if len(middle) <= 2:
        return middle[0], middle[-1] if len(middle) > 1 else ""
    for i, part in enumerate(middle[:-1]):
        if re.search(r"[.!?»\"]$", part.strip()):
            return ";".join(middle[: i + 1]), "; ".join(p for p in middle[i + 1 :] if p)
    return ";".join(middle[:-1]), middle[-1]  # une définition peut contenir des ';'


def level_from_tag(tag: str) -> str | None:
    m = re.search(r"_(Easy|Medium|Hard)$", tag)
    if not m:
        return None
    return {"Easy": "debutant", "Medium": "intermediaire", "Hard": "avance"}[m.group(1)]


def topic_from_tag(tag: str) -> str:
    base = re.sub(r"_(Easy|Medium|Hard)$", "", tag.lstrip("#"))
    return slugify(base.split("_")[0])


def subject_from_tag(tag: str) -> str:
    for prefix, subject in SUBJECT_BY_TAG_PREFIX.items():
        if tag.startswith(prefix):
            return subject
    return "marketing"


def python_topic(term: str) -> str:
    for topic, terms in PYTHON_TOPICS.items():
        if terms and term in terms:
            return topic
    if term.startswith(("pd.", "df.", "df[", "import pandas")):
        return "pandas"
    if term.endswith("Error"):
        return "erreurs"
    if term.startswith("."):
        return "chaines"
    return "bases"


def main(src: str, dst: str) -> int:
    wb = openpyxl.load_workbook(src, read_only=True)
    cards: "OrderedDict[str, dict]" = OrderedDict()
    warnings: list[str] = []

    def add(term: str, definition: str, example: str, subject: str, topic: str,
            level: str | None, tags: list[str], sheet: str) -> None:
        term = term.strip()
        key = term.lower()
        if not term or not definition:
            warnings.append(f"{sheet}: ligne vide ou sans définition pour « {term} »")
            return
        if key in cards:
            # Sheet1 fait foi ; les autres feuilles complètent seulement l'exemple manquant.
            existing = cards[key]
            if not existing["example"] and example:
                existing["example"] = example
            existing["sources"].append(sheet)
            return
        cards[key] = {
            "term": term,
            "definition": definition,
            "example": example,
            "subject": subject,
            "topic": topic,
            "level": level,
            "tags": tags,
            "sources": [sheet],
        }

    # --- Sheet1 : terme;définition;exemple;#tag -------------------------------
    for (cell,) in wb["Sheet1"].iter_rows(values_only=True):
        if not cell:
            continue
        parts = [p.strip() for p in str(cell).split(";")]
        if len(parts) < 4:
            warnings.append(f"Sheet1: ligne inexploitable : {cell[:60]}…")
            continue
        term, tag = parts[0], parts[-1]
        definition, example = split_sheet1(parts[1:-1])
        add(term, clean(definition), clean(example), subject_from_tag(tag),
            topic_from_tag(tag), level_from_tag(tag), [tag], "Sheet1")

    # --- V2 : marketing, deux colonnes ---------------------------------------
    for term, cell in wb["V2"].iter_rows(values_only=True):
        if not term or not cell:
            continue
        definition, example = split_v2(str(cell))
        add(str(term), definition, example, "marketing", "general", None, [], "V2")

    # --- Code + Code V2 : python ---------------------------------------------
    for sheet in ("Code", "Code V2"):
        for term, cell in wb[sheet].iter_rows(values_only=True):
            if not term or not cell:
                continue
            definition, example = split_v2(str(cell))
            add(str(term), definition, example, "python", python_topic(str(term)), None, [], sheet)

    # --- Identifiants, trous, audit ------------------------------------------
    out = []
    seen_ids: set[str] = set()
    for card in cards.values():
        base = f"{card['subject']}-{slugify(card['term'])}"
        cid, n = base, 2
        while cid in seen_ids:
            cid, n = f"{base}-{n}", n + 1
        seen_ids.add(cid)

        cloze = CLOZE.findall(card["definition"])
        entry = {
            "id": cid,
            "subject": card["subject"],
            "topic": card["topic"],
            "level": card["level"],
            "term": card["term"],
            "definition": CLOZE.sub(r"\1", card["definition"]),
            "example": card["example"],
            "tags": card["tags"],
        }
        if cloze:
            entry["cloze"] = card["definition"]  # version avec les {{c1::…}} intacts
        if len(entry["definition"]) > 260:
            warnings.append(f"{cid}: définition longue ({len(entry['definition'])} car.)")
        out.append(entry)

    payload = json.dumps({"version": 1, "cards": out}, ensure_ascii=False, indent=1)
    digest = hashlib.sha256(payload.encode()).hexdigest()[:8]
    with open(dst, "w", encoding="utf-8") as fh:
        fh.write(payload + "\n")

    by_subject = Counter(c["subject"] for c in out)
    by_topic = Counter(f"{c['subject']}/{c['topic']}" for c in out)
    print(f"{len(out)} cartes écrites dans {dst} (hash {digest})")
    print("Par matière :", dict(by_subject))
    print("Par sous-thème :")
    for k, v in sorted(by_topic.items()):
        print(f"  {k}: {v}")
    print(f"Cartes à trous : {sum(1 for c in out if 'cloze' in c)}")
    print(f"Sans exemple : {sum(1 for c in out if not c['example'])}")
    if warnings:
        print(f"\n{len(warnings)} avertissement(s) :")
        for w in warnings[:30]:
            print("  -", w)
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
